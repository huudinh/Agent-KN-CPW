# -*- coding: utf-8 -*-
"""
get-brief.py — Trích brief viết bài của MỘT cụm chủ đề từ bộ truy vấn GEO.

    python tools/get-brief.py mui C11
    python tools/get-brief.py Sacto C06 --core          # chỉ truy vấn CORE 80/20
    python tools/get-brief.py mui --list                # liệt kê 15 cụm của bộ
    python tools/get-brief.py --sets                    # liệt kê 17 bộ dịch vụ
    python tools/get-brief.py mui C11 -o brief-C11.md   # ghi ra file

Đầu ra là brief Markdown dán thẳng vào Agent: hành vi khách của giai đoạn,
danh sách truy vấn khách gõ, góc hỏi, khối trả lời cần có, từ khoá, rào pháp lý.

Nguồn: ../Agent-KN-GEO-Plant/doc/BotruyvanGEO/<bộ>.xlsx
"""
import argparse
import glob
import os
import sys
from collections import Counter

import openpyxl

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.normpath(os.path.join(HERE, "..", "Agent-KN-GEO-Plant", "doc", "BotruyvanGEO"))

# sheet 3 — chỉ số cột (1-based)
C_ID, C_CUM, C_GD, C_TRUYVAN, C_GOC = 1, 2, 3, 4, 5
C_NHIET, C_UUTIEN, C_GIATRI, C_NHOM, C_TANG = 6, 7, 8, 9, 10
C_DIEM, C_TUKHOA, C_KHOITL, C_TRANGDICH, C_RAO = 12, 13, 14, 15, 16


def g(ws, r, c):
    v = ws.cell(r, c).value
    return "" if v is None else str(v).strip()


def open_set(key):
    path = os.path.join(SRC, key + ".xlsx")
    if not os.path.exists(path):
        hits = [p for p in glob.glob(os.path.join(SRC, "*.xlsx"))
                if key.lower() in os.path.basename(p).lower()
                and not os.path.basename(p).startswith("~")]
        if len(hits) != 1:
            sys.exit(f"Không tìm thấy bộ '{key}'. Dùng --sets để xem danh sách.")
        path = hits[0]
    return openpyxl.load_workbook(path, data_only=True), os.path.basename(path)


def sheet(wb, prefix):
    return wb[[n for n in wb.sheetnames if n.strip().startswith(prefix)][0]]


def list_sets():
    print("17 bộ dịch vụ:\n")
    for p in sorted(glob.glob(os.path.join(SRC, "*.xlsx"))):
        b = os.path.basename(p)
        if b.startswith("~") or b.startswith("BOTRUYVAN"):
            continue
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
        ws = sheet(wb, "3.")
        print(f"  {os.path.splitext(b)[0]:<12} {g(ws, 2, 1)[:70]}")
        wb.close()


def list_clusters(key):
    wb, fname = open_set(key)
    ws = sheet(wb, "4.")
    print(f"Các cụm của bộ {fname}:\n")
    print(f"  {'MÃ':<5} {'TẦNG':<5} {'GĐ':<4} {'BÀI':>4}  CỤM")
    for r in range(6, ws.max_row + 1):
        code = g(ws, r, 2)
        if not code.startswith("C"):
            continue
        print(f"  {code:<5} {g(ws, r, 4)[:1]:<5} {g(ws, r, 5)[:2]:<4} "
              f"{_int(g(ws, r, 10)) or '—':>4}  {g(ws, r, 3)}")
    wb.close()


def build_brief(key, code, core_only):
    wb, fname = open_set(key)
    ws3, ws2, ws4 = sheet(wb, "3."), sheet(wb, "2."), sheet(wb, "4.")

    rows = [r for r in range(5, ws3.max_row + 1)
            if g(ws3, r, C_CUM).startswith(code + " ")
            or g(ws3, r, C_CUM).split("·")[0].strip() == code]
    if core_only:
        rows = [r for r in rows if "CORE" in g(ws3, r, C_NHOM)]
    if not rows:
        sys.exit(f"Không có truy vấn nào cho cụm '{code}' trong {fname}. Dùng --list để xem cụm.")

    cum = g(ws3, rows[0], C_CUM)
    gd = g(ws3, rows[0], C_GD)

    # dòng bản đồ nội dung của cụm
    mapping = {}
    for r in range(6, ws4.max_row + 1):
        if g(ws4, r, 2) == code:
            mapping = dict(tang=g(ws4, r, 4), tv=g(ws4, r, 6), core=g(ws4, r, 7),
                           art=str(_int(g(ws4, r, 10)) or "—"), need=g(ws4, r, 11),
                           url=g(ws4, r, 12), wave=g(ws4, r, 13), owner=g(ws4, r, 14),
                           role=g(ws4, r, 9))
            break

    # hành vi khách của giai đoạn
    beh = {}
    for r in range(5, 11):
        if g(ws2, r, 1) and g(ws2, r, 1).split("·")[0].strip() == gd.split("·")[0].strip():
            beh = dict(think=g(ws2, r, 3), fear=g(ws2, r, 4), want=g(ws2, r, 5),
                       trigger=g(ws2, r, 6), must=g(ws2, r, 7))
            break

    L = []
    w = L.append
    w(f"# BRIEF VIẾT BÀI — {cum}")
    w("")
    w(f"**Bộ:** `{fname}` · **Giai đoạn:** {gd} · **Tầng:** {mapping.get('tang', '—')} "
      f"· **Sóng:** {mapping.get('wave', '—')}")
    w(f"**Khối lượng:** {mapping.get('art', '—')} bài · "
      f"{mapping.get('tv', '—')} truy vấn ({mapping.get('core', '—')} CORE) "
      f"· **Trang đích:** {mapping.get('url', '—')} · **Phụ trách:** {mapping.get('owner', '—')}")
    w("")
    if mapping.get("role"):
        w(f"**Vai trò trong phễu:** {mapping['role']}")
        w("")
    if mapping.get("need"):
        w(f"**Dạng nội dung cần sản xuất:** {mapping['need']}")
        w("")

    if beh:
        w("## 1 · KHÁCH Ở GIAI ĐOẠN NÀY — ĐỌC TRƯỚC KHI VIẾT")
        w("")
        w(f"- **Đang nghĩ gì:** {beh['think']}")
        w(f"- **NỖI SỢ CHI PHỐI:** {beh['fear']}")
        w(f"- **Khát khao thật:** {beh['want']}")
        w(f"- **Trigger chuyển giai đoạn:** {beh['trigger']}")
        w(f"- **Kangnam phải có gì:** {beh['must']}")
        w("")
        w("> Nỗi sợ trên phải được gọi tên trong **150 từ đầu**, bằng chính lời khách.")
        w("")

    angles = Counter(g(ws3, r, C_GOC) for r in rows if g(ws3, r, C_GOC))
    w("## 2 · GÓC HỎI CẦN PHỦ")
    w("")
    w(" · ".join(f"**{k}** ({v})" for k, v in angles.most_common()))
    w("")

    raos = Counter(g(ws3, r, C_RAO).split("·")[0].strip() for r in rows if g(ws3, r, C_RAO))
    w("## 3 · RÀO PHÁP LÝ PHẢI XỬ LÝ")
    w("")
    seen = set()
    for r in rows:
        rao = g(ws3, r, C_RAO)
        ma = rao.split("·")[0].strip()
        if rao and ma not in seen:
            seen.add(ma)
            w(f"- **{rao}** — {raos[ma]} truy vấn")
    w("")

    w(f"## 4 · TRUY VẤN & KHỐI TRẢ LỜI CẦN CÓ ({len(rows)} truy vấn)")
    w("")
    w("Mỗi dòng là một H2/H3 tiềm năng. Dùng **nguyên văn truy vấn** làm heading khi đọc tự nhiên.")
    w("")
    rows.sort(key=lambda r: (0 if "CORE" in g(ws3, r, C_NHOM) else 1, -_int(g(ws3, r, C_DIEM))))
    for i, r in enumerate(rows, 1):
        tag = "🔥 CORE" if "CORE" in g(ws3, r, C_NHOM) else "—"
        w(f"### {i}. {g(ws3, r, C_TRUYVAN)}")
        w("")
        w(f"`{g(ws3, r, C_ID)}` · {tag} · góc hỏi **{g(ws3, r, C_GOC)}** "
          f"· {g(ws3, r, C_NHIET)} · {g(ws3, r, C_UUTIEN)} · điểm {g(ws3, r, C_DIEM)}")
        w("")
        if g(ws3, r, C_KHOITL):
            w(f"> **Khối trả lời cần có:** {g(ws3, r, C_KHOITL)}")
            w("")
        meta = []
        if g(ws3, r, C_TUKHOA):
            meta.append(f"Từ khoá: `{g(ws3, r, C_TUKHOA)}`")
        if g(ws3, r, C_TRANGDICH):
            meta.append(f"Trang đích: {g(ws3, r, C_TRANGDICH)}")
        if meta:
            w(" · ".join(meta))
            w("")

    wb.close()
    return "\n".join(L)


def _int(s):
    try:
        return int(float(s))
    except (TypeError, ValueError):
        return 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("bo", nargs="?", help="tên file bộ dịch vụ, bỏ đuôi .xlsx (vd: mui)")
    ap.add_argument("cum", nargs="?", help="mã cụm (vd: C11)")
    ap.add_argument("--list", action="store_true", help="liệt kê các cụm của bộ")
    ap.add_argument("--sets", action="store_true", help="liệt kê 17 bộ dịch vụ")
    ap.add_argument("--core", action="store_true", help="chỉ lấy truy vấn CORE 80/20")
    ap.add_argument("-o", "--out", help="ghi brief ra file")
    a = ap.parse_args()

    if a.sets:
        return list_sets()
    if not a.bo:
        ap.error("thiếu tên bộ dịch vụ (hoặc dùng --sets)")
    if a.list or not a.cum:
        return list_clusters(a.bo)

    text = build_brief(a.bo, a.cum.upper(), a.core)
    if a.out:
        with open(a.out, "w", encoding="utf-8") as fh:
            fh.write(text)
        print(f"✔ {a.out}")
    else:
        print(text)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
